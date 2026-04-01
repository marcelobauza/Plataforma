"""Hoja 3: Análisis de Sensibilidad"""
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

BLUE_HDR = PatternFill("solid", fgColor="1F3864")
BLUE_SUB = PatternFill("solid", fgColor="2E75B6")
GREEN_BG = PatternFill("solid", fgColor="E2EFDA")
ORANGE   = PatternFill("solid", fgColor="FCE4D6")
YELLOW   = PatternFill("solid", fgColor="FFFF99")
GRAY     = PatternFill("solid", fgColor="F2F2F2")

FONT_HDR  = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
FONT_BOLD = Font(name="Calibri", bold=True, size=10)
FONT_NORM = Font(name="Calibri", size=10)
FONT_SMALL= Font(name="Calibri", size=9, color="595959", italic=True)

ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_R = Alignment(horizontal="right",  vertical="center")
ALIGN_L = Alignment(horizontal="left",   vertical="center")

BORDER_ALL = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin")
)
BORDER_MED = Border(
    left=Side(style="medium"), right=Side(style="medium"),
    top=Side(style="medium"),  bottom=Side(style="medium")
)

P = "Parámetros"
F = "Flujo del Fondo"

def col(c): return get_column_letter(c)

R = {
    "n_familias":      5,
    "cuota_mensual":   6,
    "tasa_desercion":  7,
    "precio_terreno":  10,
    "viv_x_terreno":   12,
    "apreciacion":     13,
    "cap2_monto":      15,
    "cap2_ratio":      16,
    "cap2_retorno":    17,
    "gracia_meses":    19,
    "subsidio_uf":     21,
    "valor_uf":        22,
    "pct_subsidio":    23,
    "subsidio_destino":24,
    "canon_mensual":   26,
    "canon_crec":      27,
}
def p(key): return f"'{P}'!$B${R[key]}"

def crear_sensibilidad(wb):
    ws = wb.create_sheet("Análisis de Sensibilidad")
    ws.sheet_view.showGridLines = False

    # Título principal
    ws.merge_cells("A1:P1")
    t = ws["A1"]
    t.value     = "ANÁLISIS DE SENSIBILIDAD — Tablas Bidimensionales"
    t.font      = Font(name="Calibri", bold=True, color="1F3864", size=13)
    t.alignment = ALIGN_C
    t.fill      = PatternFill("solid", fgColor="D6E4F7")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:P2")
    ws["A2"].value     = ("Nota: Las tablas de sensibilidad muestran cálculos directos basados en las fórmulas del modelo. "
                          "Cambiar los parámetros en la hoja 'Parámetros' actualiza todos los resultados.")
    ws["A2"].font      = FONT_SMALL
    ws["A2"].alignment = ALIGN_L
    ws["A2"].fill      = PatternFill("solid", fgColor="FFF2CC")
    ws.row_dimensions[2].height = 18

    # ─────────────────────────────────────────────────────────────────
    # TABLA 1: Terrenos adquiridos a 5 años (mes 60)
    # Eje X: Cuota mensual [20000, 25000, 30000, 35000, 40000]
    # Eje Y: Nro familias  [100, 200, 300, 500, 1000]
    # ─────────────────────────────────────────────────────────────────
    START_ROW_T1 = 4
    cuotas   = [20000, 25000, 30000, 35000, 40000]
    familias = [100, 200, 300, 500, 1000]

    ws.merge_cells(f"A{START_ROW_T1}:G{START_ROW_T1}")
    h = ws[f"A{START_ROW_T1}"]
    h.value     = "TABLA 1 — Terrenos Adquiridos a 5 Años (mes 60)"
    h.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    h.fill      = BLUE_HDR
    h.alignment = ALIGN_C
    h.border    = BORDER_ALL
    ws.row_dimensions[START_ROW_T1].height = 22

    # Fila eje X header
    r = START_ROW_T1 + 1
    ws.cell(row=r, column=1).value     = "Familias \\ Cuota"
    ws.cell(row=r, column=1).font      = FONT_BOLD
    ws.cell(row=r, column=1).fill      = BLUE_SUB
    ws.cell(row=r, column=1).alignment = ALIGN_C
    ws.cell(row=r, column=1).border    = BORDER_ALL
    for ci, cuota in enumerate(cuotas, 2):
        ws.cell(row=r, column=ci).value          = cuota
        ws.cell(row=r, column=ci).font           = FONT_HDR
        ws.cell(row=r, column=ci).fill           = BLUE_SUB
        ws.cell(row=r, column=ci).alignment      = ALIGN_C
        ws.cell(row=r, column=ci).number_format  = '#,##0'
        ws.cell(row=r, column=ci).border         = BORDER_ALL
    ws.row_dimensions[r].height = 20

    # Filas de familias y cálculos
    for ri, fam in enumerate(familias):
        row = r + 1 + ri
        ws.cell(row=row, column=1).value          = fam
        ws.cell(row=row, column=1).font           = FONT_BOLD
        ws.cell(row=row, column=1).fill           = BLUE_SUB
        ws.cell(row=row, column=1).alignment      = ALIGN_C
        ws.cell(row=row, column=1).number_format  = '#,##0'
        ws.cell(row=row, column=1).border         = BORDER_ALL

        for ci, cuota in enumerate(cuotas, 2):
            # Fórmula directa de terrenos a mes 60:
            # Acum C1 en mes 60 ≈ suma de aportes mensuales con deserción
            # Aprox: aporte_mes_t = fam * (1-des)^(t/12) * cuota
            # Acum_C1_60 ≈ cuota * fam * sum_{t=1}^{60} (1-des)^((t-1)/12)
            # Des mensual = (1-des_anual)^(1/12)
            # sum geométrico: fam * cuota * (1 - des_mes^60) / (1 - des_mes)
            # Acum_C2_60 ≈ min(cap2_monto, acum_c1 * ratio)
            # Cap_total = Acum_C1 + Acum_C2 + subsidio
            # Terrenos = ENTERO(cap_total / precio*(1+aprec)^(60/12))
            f = (
                f"=LET("
                f"des_mes,(1-{p('tasa_desercion')})^(1/12),"
                f"acum_c1,{cuota}*{fam}*(1-des_mes^60)/(1-des_mes),"
                f"acum_c2,MIN({p('cap2_monto')},acum_c1*{p('cap2_ratio')}),"
                f"sub_por_t,{p('viv_x_terreno')}*{p('pct_subsidio')}*{p('subsidio_uf')}*{p('valor_uf')}*{p('subsidio_destino')},"
                f"cap_total,acum_c1+acum_c2,"
                f"precio_adj,{p('precio_terreno')}*(1+{p('apreciacion')})^(60/12),"
                f"terrenos_aprox,MAX(0,ENTERO(cap_total/precio_adj)),"
                f"ENTERO((cap_total+terrenos_aprox*sub_por_t)/precio_adj)"
                f")"
            )
            ws.cell(row=row, column=ci).value         = f
            ws.cell(row=row, column=ci).number_format = '0'
            ws.cell(row=row, column=ci).font          = FONT_NORM
            ws.cell(row=row, column=ci).alignment     = ALIGN_C
            ws.cell(row=row, column=ci).border        = BORDER_ALL

        ws.row_dimensions[row].height = 18

    # Escala de color Tabla 1
    t1_data = f"B{r+1}:F{r+len(familias)}"
    ws.conditional_formatting.add(t1_data, ColorScaleRule(
        start_type='min', start_color='FCE4D6',
        mid_type='percentile', mid_value=50, mid_color='FFFF99',
        end_type='max', end_color='E2EFDA'
    ))

    # ─────────────────────────────────────────────────────────────────
    # TABLA 2: TIR Capa 2 (proxy: canon / servicio deuda a 15 años)
    # Eje X: Canon mensual  [15000, 20000, 25000, 30000, 35000]
    # Eje Y: Retorno Capa 2 [3%, 4%, 5%, 6%, 7%]
    # ─────────────────────────────────────────────────────────────────
    START_ROW_T2 = r + len(familias) + 3
    canones  = [15000, 20000, 25000, 30000, 35000]
    retornos = [0.03, 0.04, 0.05, 0.06, 0.07]

    ws.merge_cells(f"A{START_ROW_T2}:G{START_ROW_T2}")
    h2 = ws[f"A{START_ROW_T2}"]
    h2.value     = "TABLA 2 — Ratio Cobertura Deuda Capa 2 a 15 Años (Canon / Servicio Deuda)"
    h2.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    h2.fill      = PatternFill("solid", fgColor="375623")
    h2.alignment = ALIGN_C
    h2.border    = BORDER_ALL
    ws.row_dimensions[START_ROW_T2].height = 22

    r2 = START_ROW_T2 + 1
    ws.cell(row=r2, column=1).value     = "Ret.C2 \\ Canon"
    ws.cell(row=r2, column=1).font      = FONT_BOLD
    ws.cell(row=r2, column=1).fill      = PatternFill("solid", fgColor="548235")
    ws.cell(row=r2, column=1).alignment = ALIGN_C
    ws.cell(row=r2, column=1).border    = BORDER_ALL
    for ci, can in enumerate(canones, 2):
        ws.cell(row=r2, column=ci).value         = can
        ws.cell(row=r2, column=ci).font          = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        ws.cell(row=r2, column=ci).fill          = PatternFill("solid", fgColor="548235")
        ws.cell(row=r2, column=ci).alignment     = ALIGN_C
        ws.cell(row=r2, column=ci).number_format = '#,##0'
        ws.cell(row=r2, column=ci).border        = BORDER_ALL
    ws.row_dimensions[r2].height = 20

    for ri, ret in enumerate(retornos):
        row2 = r2 + 1 + ri
        ws.cell(row=row2, column=1).value         = ret
        ws.cell(row=row2, column=1).font          = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        ws.cell(row=row2, column=1).fill          = PatternFill("solid", fgColor="548235")
        ws.cell(row=row2, column=1).alignment     = ALIGN_C
        ws.cell(row=row2, column=1).number_format = '0.0%'
        ws.cell(row=row2, column=1).border        = BORDER_ALL

        for ci, can in enumerate(canones, 2):
            # Ratio cobertura = Canon total año 15 / Servicio deuda año 15
            # Canon año 15: familias_ubicadas_180 * canon_ajustado
            # Servicio deuda: cap2_desembolsado * retorno / 12 * 12
            # Usando cap2_monto como proxy del capital desembolsado total
            f2 = (
                f"=LET("
                f"des_mes,(1-{p('tasa_desercion')})^(1/12),"
                f"acum_c1,{p('cuota_mensual')}*{p('n_familias')}*(1-des_mes^180)/(1-des_mes),"
                f"acum_c2,MIN({p('cap2_monto')},acum_c1*{p('cap2_ratio')}),"
                f"sub_por_t,{p('viv_x_terreno')}*{p('pct_subsidio')}*{p('subsidio_uf')}*{p('valor_uf')}*{p('subsidio_destino')},"
                f"cap_total,acum_c1+acum_c2,"
                f"precio_adj,{p('precio_terreno')}*(1+{p('apreciacion')})^(180/12),"
                f"terrenos,MAX(0,ENTERO(cap_total/precio_adj)),"
                f"fam_ubic,terrenos*{p('viv_x_terreno')},"
                f"canon_anual,fam_ubic*{can}*(1+{p('canon_crec')})^(180/12)*12,"
                f"serv_deuda,acum_c2*{ret},"
                f"IF(serv_deuda=0,999,REDONDEAR(canon_anual/serv_deuda,2))"
                f")"
            )
            ws.cell(row=row2, column=ci).value         = f2
            ws.cell(row=row2, column=ci).number_format = '0.00"x"'
            ws.cell(row=row2, column=ci).font          = FONT_NORM
            ws.cell(row=row2, column=ci).alignment     = ALIGN_C
            ws.cell(row=row2, column=ci).border        = BORDER_ALL

        ws.row_dimensions[row2].height = 18

    ws.conditional_formatting.add(f"B{r2+1}:F{r2+len(retornos)}", ColorScaleRule(
        start_type='min', start_color='FCE4D6',
        mid_type='num', mid_value=1, mid_color='FFFF99',
        end_type='max', end_color='E2EFDA'
    ))

    # ─────────────────────────────────────────────────────────────────
    # TABLA 3: Meses hasta primer sorteo
    # Eje X: Precio terreno [80M, 100M, 120M, 150M, 200M]
    # Eje Y: Nro familias   [100, 200, 300, 500, 1000]
    # ─────────────────────────────────────────────────────────────────
    START_ROW_T3 = row2 + len(retornos) + 3
    precios  = [80000000, 100000000, 120000000, 150000000, 200000000]
    familias3= [100, 200, 300, 500, 1000]

    ws.merge_cells(f"A{START_ROW_T3}:G{START_ROW_T3}")
    h3 = ws[f"A{START_ROW_T3}"]
    h3.value     = "TABLA 3 — Meses Hasta Primer Sorteo (primer terreno adquirido)"
    h3.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    h3.fill      = PatternFill("solid", fgColor="7030A0")
    h3.alignment = ALIGN_C
    h3.border    = BORDER_ALL
    ws.row_dimensions[START_ROW_T3].height = 22

    r3 = START_ROW_T3 + 1
    ws.cell(row=r3, column=1).value     = "Familias \\ Precio"
    ws.cell(row=r3, column=1).font      = FONT_BOLD
    ws.cell(row=r3, column=1).fill      = PatternFill("solid", fgColor="9B59B6")
    ws.cell(row=r3, column=1).alignment = ALIGN_C
    ws.cell(row=r3, column=1).border    = BORDER_ALL
    for ci, prec in enumerate(precios, 2):
        ws.cell(row=r3, column=ci).value         = prec
        ws.cell(row=r3, column=ci).font          = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        ws.cell(row=r3, column=ci).fill          = PatternFill("solid", fgColor="9B59B6")
        ws.cell(row=r3, column=ci).alignment     = ALIGN_C
        ws.cell(row=r3, column=ci).number_format = '#,##0'
        ws.cell(row=r3, column=ci).border        = BORDER_ALL
    ws.row_dimensions[r3].height = 20

    for ri, fam in enumerate(familias3):
        row3 = r3 + 1 + ri
        ws.cell(row=row3, column=1).value         = fam
        ws.cell(row=row3, column=1).font          = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        ws.cell(row=row3, column=1).fill          = PatternFill("solid", fgColor="9B59B6")
        ws.cell(row=row3, column=1).alignment     = ALIGN_C
        ws.cell(row=row3, column=1).number_format = '#,##0'
        ws.cell(row=row3, column=1).border        = BORDER_ALL

        for ci, prec in enumerate(precios, 2):
            # Mes donde capital total >= precio terreno (ajustado por apreciación)
            # Capital mes t ≈ cuota * fam * (1-des^t)/(1-des) * (1 + ratio)
            # Resolución numérica: iteramos de 1 a 180
            # Usamos fórmula MATCH sobre array implícito
            f3 = (
                f"=LET("
                f"des_mes,(1-{p('tasa_desercion')})^(1/12),"
                f"meses,SECUENCIA(180),"
                f"acum_c1,{p('cuota_mensual')}*{fam}*(1-des_mes^meses)/(1-des_mes),"
                f"acum_c2,MIN({p('cap2_monto')},acum_c1*{p('cap2_ratio')}),"
                f"cap_total,acum_c1+acum_c2,"
                f"precio_adj,{prec}*(1+{p('apreciacion')})^((meses-1)/12),"
                f"puede,cap_total>=precio_adj,"
                f"idx,COINCIDIR(VERDADERO,puede,0),"
                f"IF(ESERROR(idx),\">180\",idx)"
                f")"
            )
            ws.cell(row=row3, column=ci).value         = f3
            ws.cell(row=row3, column=ci).number_format = '0'
            ws.cell(row=row3, column=ci).font          = FONT_NORM
            ws.cell(row=row3, column=ci).alignment     = ALIGN_C
            ws.cell(row=row3, column=ci).border        = BORDER_ALL

        ws.row_dimensions[row3].height = 18

    ws.conditional_formatting.add(f"B{r3+1}:F{r3+len(familias3)}", ColorScaleRule(
        start_type='min', start_color='E2EFDA',
        mid_type='percentile', mid_value=50, mid_color='FFFF99',
        end_type='max', end_color='FCE4D6'
    ))

    # Anchos de columna
    ws.column_dimensions["A"].width = 18
    for ci in range(2, 7):
        ws.column_dimensions[col(ci)].width = 16

    ws.freeze_panes = "B4"
    return ws

print("sensibilidad.py cargado OK")
